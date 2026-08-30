from datetime import date, timedelta
from io import BytesIO
from decimal import Decimal

from django.contrib import messages
from usuarios.mixins import RolRequiredMixin, SOLO_ADMIN
from django.db.models import F, Sum, Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.utils import timezone
from django.views import View
from django.views.generic import ListView, CreateView, TemplateView

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from distribucion.models import Planilla, Entrega, Credito
from produccion.models import Produccion, Insumo
from usuarios.models import Usuario
from .forms import DescuadreForm, FiltroVentasForm
from .models import Descuadre


# ── Dashboard ─────────────────────────────────────────────────────────────────

class ReportesDashboardView(RolRequiredMixin, TemplateView):
    roles_permitidos = SOLO_ADMIN
    template_name = 'reportes/dashboard.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        hoy = date.today()
        inicio_mes = hoy.replace(day=1)

        ctx['total_produccion_mes'] = (
            Produccion.objects
            .filter(fecha__gte=inicio_mes)
            .aggregate(total=Sum('cantidad_producida'))['total'] or 0
        )
        ctx['total_ventas_mes'] = sum(
            e.subtotal for e in
            Entrega.objects.filter(planilla__fecha__gte=inicio_mes).select_related('producto')
        )
        ctx['creditos_pendientes'] = (
            Credito.objects.filter(pagado=False)
            .aggregate(total=Sum('saldo_pendiente'))['total'] or 0
        )
        ctx['alertas_insumos'] = (
            Insumo.objects.filter(
                activo=True, stock_actual__lte=F('stock_minimo')
            ).count()
        )
        ctx['descuadres_activos'] = Descuadre.objects.filter(resuelto=False).count()
        ctx['planillas_pendientes'] = Planilla.objects.filter(
            estado=Planilla.Estado.PENDIENTE_VALIDACION
        ).count()
        ctx['hoy'] = str(hoy)            # ISO string para URLs
        ctx['inicio_mes'] = str(inicio_mes)  # primer día del mes
        return ctx


# ── Ventas consolidadas ───────────────────────────────────────────────────────

class VentasConsolidadasView(RolRequiredMixin, TemplateView):
    roles_permitidos = SOLO_ADMIN
    template_name = 'reportes/ventas.html'

    def _get_filtros(self):
        hoy = date.today()
        distribuidores = Usuario.objects.filter(rol='DIST')
        form = FiltroVentasForm(
            distribuidores,
            self.request.GET or {'fecha_desde': hoy.replace(day=1), 'fecha_hasta': hoy},
        )
        fecha_desde = None
        fecha_hasta = None
        distribuidor_id = None
        if form.is_valid():
            fecha_desde = form.cleaned_data.get('fecha_desde') or hoy.replace(day=1)
            fecha_hasta = form.cleaned_data.get('fecha_hasta') or hoy
            distribuidor_id = form.cleaned_data.get('distribuidor') or None
        return form, fecha_desde, fecha_hasta, distribuidor_id

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        form, f_desde, f_hasta, dist_id = self._get_filtros()
        ctx['form'] = form

        qs = Entrega.objects.filter(
            planilla__fecha__gte=f_desde,
            planilla__fecha__lte=f_hasta,
        ).select_related('producto', 'cliente', 'planilla__distribuidor')

        if dist_id:
            qs = qs.filter(planilla__distribuidor_id=dist_id)

        # Ventas por producto
        por_producto = {}
        for e in qs:
            key = e.producto.nombre
            if key not in por_producto:
                por_producto[key] = {'cantidad': 0, 'total': 0}
            por_producto[key]['cantidad'] += e.cantidad - e.devolucion
            por_producto[key]['total'] += float(e.subtotal)

        # Ventas por distribuidor
        por_dist = {}
        for e in qs:
            nombre = e.planilla.distribuidor.get_full_name() or e.planilla.distribuidor.username
            if nombre not in por_dist:
                por_dist[nombre] = {'entregas': 0, 'total': 0}
            por_dist[nombre]['entregas'] += 1
            por_dist[nombre]['total'] += float(e.subtotal)

        # Desglose por modalidad de pago
        pago_totales = {'EFE': 0, 'NEQ': 0, 'CRE': 0}
        for e in qs:
            pago_totales[e.modalidad_pago] = pago_totales.get(e.modalidad_pago, 0) + float(e.subtotal)

        ctx['por_producto'] = sorted(por_producto.items(), key=lambda x: -x[1]['total'])
        ctx['por_dist'] = sorted(por_dist.items(), key=lambda x: -x[1]['total'])
        ctx['pago_totales'] = pago_totales
        ctx['gran_total'] = sum(v['total'] for v in por_producto.values())
        ctx['fecha_desde'] = str(f_desde)   # ISO YYYY-MM-DD para URLs y date inputs
        ctx['fecha_hasta'] = str(f_hasta)
        return ctx


class ExportarVentasExcelView(RolRequiredMixin, View):
    roles_permitidos = SOLO_ADMIN
    def get(self, request):
        hoy = date.today()
        f_desde = request.GET.get('fecha_desde') or str(hoy.replace(day=1))
        f_hasta = request.GET.get('fecha_hasta') or str(hoy)

        qs = Entrega.objects.filter(
            planilla__fecha__gte=f_desde,
            planilla__fecha__lte=f_hasta,
        ).select_related('producto', 'cliente', 'planilla__distribuidor').order_by('planilla__fecha')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ventas'

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='1F4E79')
        headers = ['Fecha', 'Distribuidor', 'Cliente', 'Producto',
                   'Cantidad', 'Devolución', 'Precio unit.', 'Subtotal', 'Pago']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row, e in enumerate(qs, 2):
            ws.append([
                e.planilla.fecha,
                e.planilla.distribuidor.get_full_name() or e.planilla.distribuidor.username,
                e.cliente.nombre,
                e.producto.nombre,
                e.cantidad,
                e.devolucion,
                float(e.precio_unitario),
                float(e.subtotal),
                e.get_modalidad_pago_display(),
            ])

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or '')) for c in col) + 4

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        nombre = f'ventas_{f_desde}_{f_hasta}.xlsx'
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{nombre}"'
        return resp


# ── Descuadres ────────────────────────────────────────────────────────────────

class DescuadreListView(RolRequiredMixin, ListView):
    roles_permitidos = SOLO_ADMIN
    model = Descuadre
    template_name = 'reportes/descuadres.html'
    context_object_name = 'descuadres'
    paginate_by = 10

    def get_queryset(self):
        qs = Descuadre.objects.select_related('detectado_por')
        if q := self.request.GET.get('q'):
            qs = qs.filter(descripcion__icontains=q)
        if tipo := self.request.GET.get('tipo'):
            qs = qs.filter(tipo=tipo)
        resuelto = self.request.GET.get('resuelto')
        if resuelto in ('1', '0'):
            qs = qs.filter(resuelto=resuelto == '1')
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['tipos'] = Descuadre.TipoDescuadre.choices
        ctx['form'] = DescuadreForm()
        return ctx


class DescuadreCreateView(RolRequiredMixin, CreateView):
    roles_permitidos = SOLO_ADMIN
    model = Descuadre
    form_class = DescuadreForm
    template_name = 'reportes/descuadres.html'

    def form_valid(self, form):
        form.instance.detectado_por = self.request.user
        descuadre = form.save()
        messages.success(self.request, f'Descuadre registrado: {descuadre}')
        return redirect('reportes:descuadres')

    def form_invalid(self, form):
        messages.error(self.request, 'Error al registrar el descuadre.')
        return redirect('reportes:descuadres')


class DescuadreResolverView(RolRequiredMixin, View):
    roles_permitidos = SOLO_ADMIN
    def post(self, request, pk):
        descuadre = get_object_or_404(Descuadre, pk=pk)
        if not descuadre.resuelto:
            descuadre.resuelto = True
            descuadre.resuelto_en = timezone.now()
            descuadre.save(update_fields=['resuelto', 'resuelto_en'])
            messages.success(request, 'Descuadre marcado como resuelto.')
        return redirect('reportes:descuadres')


# ── Créditos pendientes ───────────────────────────────────────────────────────

class CreditosPendientesView(RolRequiredMixin, ListView):
    roles_permitidos = SOLO_ADMIN
    template_name = 'reportes/creditos.html'
    context_object_name = 'creditos'
    paginate_by = 10

    def get_queryset(self):
        qs = Credito.objects.filter(pagado=False).select_related(
            'cliente', 'entrega__planilla__distribuidor', 'entrega__producto'
        )
        if cliente := self.request.GET.get('cliente'):
            qs = qs.filter(cliente__nombre__icontains=cliente)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['total_pendiente'] = (
            Credito.objects.filter(pagado=False)
            .aggregate(t=Sum('saldo_pendiente'))['t'] or 0
        )
        return ctx


class CreditoPagarView(RolRequiredMixin, View):
    roles_permitidos = SOLO_ADMIN

    def post(self, request, pk):
        credito = get_object_or_404(Credito, pk=pk, pagado=False)
        try:
            monto_pago = round(float(request.POST.get('monto_pago', 0)), 2)
        except (ValueError, TypeError):
            monto_pago = 0

        if monto_pago <= 0:
            messages.error(request, 'El monto del pago debe ser mayor a cero.')
            return redirect('reportes:creditos')

        if monto_pago > float(credito.saldo_pendiente):
            messages.error(
                request,
                f'El monto ingresado (${monto_pago}) supera el saldo pendiente '
                f'(${credito.saldo_pendiente}).',
            )
            return redirect('reportes:creditos')

        credito.saldo_pendiente = round(float(credito.saldo_pendiente) - monto_pago, 2)
        if credito.saldo_pendiente <= 0:
            credito.saldo_pendiente = 0
            credito.pagado = True
            credito.fecha_pago = date.today()
            messages.success(
                request,
                f'Crédito de {credito.cliente.nombre} pagado completamente.',
            )
        else:
            messages.success(
                request,
                f'Abono de ${monto_pago} registrado. Saldo restante: ${credito.saldo_pendiente}.',
            )
        credito.save(update_fields=['saldo_pendiente', 'pagado', 'fecha_pago'])
        return redirect('reportes:creditos')


# ── PDF export ventas ─────────────────────────────────────────────────────────

class ExportarVentasPDFView(RolRequiredMixin, View):
    roles_permitidos = SOLO_ADMIN

    def get(self, request):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, HRFlowable,
        )

        hoy = date.today()
        f_desde = request.GET.get('fecha_desde') or str(hoy.replace(day=1))
        f_hasta = request.GET.get('fecha_hasta') or str(hoy)

        qs = (
            Entrega.objects
            .filter(planilla__fecha__gte=f_desde, planilla__fecha__lte=f_hasta)
            .select_related('producto', 'cliente', 'planilla__distribuidor')
            .order_by('planilla__fecha', 'planilla__distribuidor__first_name')
        )

        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=landscape(A4),
            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
            topMargin=1.8 * cm, bottomMargin=1.5 * cm,
            title=f'Ventas {f_desde} – {f_hasta}',
        )

        NAVY = colors.HexColor('#1B2870')
        LIGHT_BLUE = colors.HexColor('#EFF3FB')
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'title', parent=styles['Heading1'],
            fontSize=14, textColor=NAVY, spaceAfter=2,
        )
        sub_style = ParagraphStyle(
            'sub', parent=styles['Normal'],
            fontSize=9, textColor=colors.HexColor('#64748b'), spaceAfter=0,
        )
        cell_style = ParagraphStyle(
            'cell', parent=styles['Normal'], fontSize=7.5,
        )

        story = []

        # Encabezado
        story.append(Paragraph('Brisas de Pacandé', title_style))
        story.append(Paragraph(
            f'Ventas consolidadas &nbsp;·&nbsp; {f_desde} al {f_hasta} &nbsp;·&nbsp; '
            f'Generado: {hoy.strftime("%d/%m/%Y")}',
            sub_style,
        ))
        story.append(HRFlowable(width='100%', thickness=1, color=NAVY, spaceAfter=10))

        # Tabla principal
        headers = ['Fecha', 'Distribuidor', 'Cliente', 'Producto',
                   'Cant.', 'Dev.', 'P. Unit.', 'Subtotal', 'Pago']
        col_widths = [2.1*cm, 3.8*cm, 4.2*cm, 4.2*cm,
                      1.5*cm, 1.4*cm, 2.2*cm, 2.4*cm, 2.2*cm]

        data = [headers]
        totales_producto = {}
        gran_total = Decimal('0')

        for e in qs:
            dist = e.planilla.distribuidor
            dist_nombre = dist.get_full_name() or dist.username
            data.append([
                e.planilla.fecha.strftime('%d/%m/%Y'),
                Paragraph(dist_nombre, cell_style),
                Paragraph(e.cliente.nombre, cell_style),
                Paragraph(e.producto.nombre, cell_style),
                str(e.cantidad),
                str(e.devolucion),
                f'${e.precio_unitario:,.0f}',
                f'${e.subtotal:,.0f}',
                e.get_modalidad_pago_display(),
            ])
            key = e.producto.nombre
            totales_producto[key] = totales_producto.get(key, Decimal('0')) + e.subtotal
            gran_total += e.subtotal

        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            # Encabezado
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            # Cuerpo
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 7.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, LIGHT_BLUE]),
            # Alineación numérica
            ('ALIGN', (4, 1), (7, -1), 'RIGHT'),
            ('ALIGN', (8, 1), (8, -1), 'CENTER'),
            # Grilla
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
            ('LINEBELOW', (0, 0), (-1, 0), 1, NAVY),
            # Padding
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.5 * cm))

        # Totales por producto
        if totales_producto:
            resumen_data = [['Producto', 'Total vendido']]
            for nombre, total in sorted(totales_producto.items(), key=lambda x: -x[1]):
                resumen_data.append([nombre, f'${total:,.0f}'])
            resumen_data.append(['GRAN TOTAL', f'${gran_total:,.0f}'])

            resumen = Table(resumen_data, colWidths=[8*cm, 3.5*cm])
            resumen.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), NAVY),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, LIGHT_BLUE]),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#DBEAFE')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(resumen)

        doc.build(story)
        buf.seek(0)
        nombre = f'ventas_{f_desde}_{f_hasta}.pdf'
        resp = HttpResponse(buf, content_type='application/pdf')
        resp['Content-Disposition'] = f'attachment; filename="{nombre}"'
        return resp
